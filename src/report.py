from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.thresholds import MULTI_PIN_THRESHOLDS, single_pin_thresholds

PCT_FORMAT = "0.00%"
AVG_FORMAT = "0.0"

GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
YELLOW_FILL = PatternFill(fill_type="solid", fgColor="FFEB9C")
RED_FILL = PatternFill(fill_type="solid", fgColor="FFC7CE")
SUMMARY_HEADER_FILL = PatternFill(fill_type="solid", fgColor="17365D")


def append_event_total_row(
    summary: pd.DataFrame,
    team_games: pd.DataFrame | None,
) -> pd.DataFrame:
    if summary.empty:
        return summary

    def summed(column: str) -> int:
        return int(pd.to_numeric(summary[column], errors="coerce").fillna(0).sum())

    single_made = summed("Single-Pin Made")
    single_attempts = summed("Single-Pin Attempts")
    multi_made = summed("Multi-Pin Made")
    multi_attempts = summed("Multi-Pin Attempts")
    strikes = summed("Strikes")
    strike_attempts = summed("Strike Attempts")
    baker_points = summed("Baker Pts")
    baker_frames = summed("Baker Frames")
    average = None
    if team_games is not None and not team_games.empty:
        scores = pd.to_numeric(team_games["score"], errors="coerce").dropna()
        average = float(scores.mean()) if not scores.empty else None

    total = {
        "Name": "Total",
        "Single-Pin Made": single_made,
        "Single-Pin Attempts": single_attempts,
        "Single-Pin %": single_made / single_attempts if single_attempts else None,
        "Multi-Pin Made": multi_made,
        "Multi-Pin Attempts": multi_attempts,
        "Multi-Pin %": multi_made / multi_attempts if multi_attempts else None,
        "Strikes": strikes,
        "Strike Attempts": strike_attempts,
        "Strike %": strikes / strike_attempts if strike_attempts else None,
        "Baker Pts": baker_points,
        "Baker Frames": baker_frames,
        "Baker Avg": baker_points / baker_frames if baker_frames else None,
        "Average": average,
    }
    return pd.concat([summary, pd.DataFrame([total])], ignore_index=True)


def add_team_game_score_columns(
    summary: pd.DataFrame,
    team_games: pd.DataFrame | None,
    athlete_filter: str | None = None,
) -> pd.DataFrame:
    if summary.empty or team_games is None or team_games.empty:
        return summary

    scores_by_row: list[list[object]] = []
    for _, summary_row in summary.iterrows():
        athlete = athlete_filter or str(summary_row["Name"])
        games = team_games[team_games["athlete"] == athlete].copy()
        if games.empty:
            scores_by_row.append([])
            continue
        games["__date"] = pd.to_datetime(games.get("game_date"), errors="coerce")
        if "game_id" in games:
            games["__game_order"] = pd.to_numeric(
                games["game_id"].astype(str).str.rsplit(":", n=1).str[-1],
                errors="coerce",
            )
        sort_columns = [
            column
            for column in ("__date", "source_file", "__game_order")
            if column in games
        ]
        games = games.sort_values(sort_columns, ascending=True, na_position="last")
        scores_by_row.append(games["score"].tolist())

    max_games = max((len(scores) for scores in scores_by_row), default=0)
    if max_games == 0:
        return summary

    out = summary.copy()
    average_column = out.pop("Average") if "Average" in out.columns else None
    for game_index in range(max_games):
        out[f"Game {game_index + 1}"] = [
            scores[game_index] if game_index < len(scores) else None
            for scores in scores_by_row
        ]
    if average_column is not None:
        out["Average"] = average_column
    return out


def build_summary_workbook(
    summary: pd.DataFrame,
    team: str,
    event: str,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overall Summary"
    ws.sheet_view.showGridLines = False

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(summary.columns), 1))
    title = ws.cell(1, 1, f"LTU {team} - {event} Overall Summary")
    title.font = Font(size=16, bold=True, color="17365D")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25

    export = summary.rename(
        columns={
            "Single-Pin Made": "SP Made",
            "Single-Pin Attempts": "SP Att",
            "Single-Pin %": "SP %",
            "Multi-Pin Made": "MP Made",
            "Multi-Pin Attempts": "MP Att",
            "Multi-Pin %": "MP %",
            "Strikes": "STR",
            "Strike Attempts": "STR Att",
            "Strike %": "STR %",
            "Baker Pts": "BKR Pts",
            "Baker Frames": "BKR Frm",
            "Baker Avg": "BKR Avg",
            "Average": "AVG",
        }
    )
    ws.append([])
    ws.append(export.columns.tolist())
    for row in export.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])

    header_row = 3
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = SUMMARY_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    single_thresholds = single_pin_thresholds(team)
    column_numbers = {cell.value: cell.column for cell in ws[header_row]}
    for row in range(header_row + 1, ws.max_row + 1):
        ws.cell(row, column_numbers["Name"]).alignment = Alignment(horizontal="left")
        if ws.cell(row, column_numbers["Name"]).value == "Total":
            for cell in ws[row]:
                cell.font = Font(bold=True)
        for header in ("SP %", "MP %", "STR %"):
            if header not in column_numbers:
                continue
            cell = ws.cell(row, column_numbers[header])
            if cell.value is not None:
                cell.number_format = PCT_FORMAT
        for header in ("BKR Avg", "AVG"):
            if header not in column_numbers:
                continue
            cell = ws.cell(row, column_numbers[header])
            if cell.value is not None:
                cell.number_format = "0.00"
        sp_cell = ws.cell(row, column_numbers["SP %"]) if "SP %" in column_numbers else None
        if sp_cell is not None and sp_cell.value is not None:
            _apply_percentage_color(
                sp_cell,
                green_threshold=single_thresholds[0],
                yellow_threshold=single_thresholds[1],
            )
        mp_cell = ws.cell(row, column_numbers["MP %"]) if "MP %" in column_numbers else None
        if mp_cell is not None and mp_cell.value is not None:
            _apply_percentage_color(
                mp_cell,
                green_threshold=MULTI_PIN_THRESHOLDS[0],
                yellow_threshold=MULTI_PIN_THRESHOLDS[1],
            )

    widths = {
        "Name": 28,
        "SP Made": 11,
        "SP Att": 10,
        "SP %": 10,
        "MP Made": 11,
        "MP Att": 10,
        "MP %": 10,
        "STR": 9,
        "STR Att": 10,
        "STR %": 10,
        "BKR Pts": 11,
        "BKR Frm": 11,
        "BKR Avg": 11,
        "AVG": 10,
    }
    for cell in ws[header_row]:
        width = 10 if str(cell.value).startswith("Game ") else widths.get(str(cell.value), 12)
        ws.column_dimensions[cell.column_letter].width = width

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(header_row, ws.max_column).column_letter}{ws.max_row}"
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_report_workbook(
    records: pd.DataFrame,
    frame_attempts: pd.DataFrame | None = None,
    baker_frames: pd.DataFrame | None = None,
    team_games: pd.DataFrame | None = None,
    team: str = "Men",
) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Overall Summary"

    single_thresholds = single_pin_thresholds(team)
    _build_summary_sheet(
        ws_summary,
        records,
        frame_attempts,
        baker_frames,
        team_games,
        single_thresholds,
    )

    athletes = _all_athletes(records, frame_attempts, baker_frames, team_games)
    used_names = {ws_summary.title}
    for athlete in athletes:
        sheet_name = _unique_sheet_name(athlete, used_names)
        used_names.add(sheet_name)
        ws = wb.create_sheet(title=sheet_name)
        _build_athlete_sheet(
            ws,
            athlete,
            records[records["athlete"] == athlete].copy(),
            single_thresholds,
        )

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _build_summary_sheet(
    ws,
    records: pd.DataFrame,
    frame_attempts: pd.DataFrame | None,
    baker_frames: pd.DataFrame | None,
    team_games: pd.DataFrame | None,
    single_thresholds: tuple[float, float],
) -> None:
    headers = [
        "Name",
        "Single-Pin Made",
        "Single-Pin Attempts",
        "Single-Pin %",
        "Multi-Pin Made",
        "Multi-Pin Attempts",
        "Multi-Pin %",
        "Strikes",
        "Strike Attempts",
        "Strike %",
        "Baker Pts",
        "Baker Frames",
        "Baker Avg",
        "Average",
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    athletes = _all_athletes(records, frame_attempts, baker_frames, team_games)
    for athlete in athletes:
        subset = records[records["athlete"] == athlete]

        multi = subset[~subset["is_single"]]
        single = subset[subset["is_single"]]

        multi_attempts = int(multi["count_attempt"].sum())
        multi_made = int((multi["made"] & multi["count_attempt"]).sum())
        single_attempts = int(single["count_attempt"].sum())
        single_made = int((single["made"] & single["count_attempt"]).sum())

        multi_pct = (multi_made / multi_attempts) if multi_attempts else None
        single_pct = (single_made / single_attempts) if single_attempts else None
        baker_points = 0
        baker_count = 0
        baker_avg = None
        if baker_frames is not None and not baker_frames.empty:
            baker_subset = baker_frames[baker_frames["athlete"] == athlete]
            baker_count = int(len(baker_subset))
            baker_points = int(baker_subset["points"].sum())
            baker_avg = (baker_points / baker_count) if baker_count else None
        strikes = 0
        strike_attempts = 0
        strike_pct = None
        if frame_attempts is not None and not frame_attempts.empty:
            strike_subset = frame_attempts[frame_attempts["athlete"] == athlete]
            strike_attempts = int(len(strike_subset))
            strikes = int(strike_subset["is_strike"].sum())
            strike_pct = (strikes / strike_attempts) if strike_attempts else None
        average = None
        if team_games is not None and not team_games.empty:
            team_subset = team_games[team_games["athlete"] == athlete]
            if not team_subset.empty:
                average = team_subset["score"].sum() / len(team_subset)

        ws.append([
            athlete,
            single_made,
            single_attempts,
            single_pct,
            multi_made,
            multi_attempts,
            multi_pct,
            strikes,
            strike_attempts,
            strike_pct,
            baker_points,
            baker_count,
            baker_avg,
            average,
        ])

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=14):
        for cell in row:
            if cell.column in (4, 7, 10) and cell.value is not None:
                cell.number_format = PCT_FORMAT
                if cell.column == 4:
                    _apply_percentage_color(
                        cell,
                        green_threshold=single_thresholds[0],
                        yellow_threshold=single_thresholds[1],
                    )
                elif cell.column == 7:
                    _apply_percentage_color(
                        cell,
                        green_threshold=MULTI_PIN_THRESHOLDS[0],
                        yellow_threshold=MULTI_PIN_THRESHOLDS[1],
                    )
            elif cell.column == 13 and cell.value is not None:
                cell.number_format = AVG_FORMAT
            elif cell.column == 14 and cell.value is not None:
                cell.number_format = "0.00"

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 14
    ws.column_dimensions["J"].width = 10
    ws.column_dimensions["K"].width = 10
    ws.column_dimensions["L"].width = 12
    ws.column_dimensions["M"].width = 10
    ws.column_dimensions["N"].width = 10


def _build_athlete_sheet(
    ws,
    athlete: str,
    records: pd.DataFrame,
    single_thresholds: tuple[float, float],
) -> None:
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = athlete
    title.font = Font(size=18, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 3
    for label, is_single in [("Single Pin", True), ("Multi Pin", False)]:
        ws.cell(row=current_row, column=1, value=label).font = Font(bold=True)
        current_row += 1
        for column, header in enumerate(["Pin Leave", "Made", "Attempts", "Make %"], start=1):
            cell = ws.cell(row=current_row, column=column, value=header)
            cell.font = Font(bold=True)
        current_row += 1

        section = _aggregate_section(records, is_single=is_single)
        common, rare, _ = _split_common_and_rare(section)

        for _, item in common.iterrows():
            _write_spare_row(
                ws,
                current_row,
                item,
                is_single=is_single,
                color_percentage=True,
                single_thresholds=single_thresholds,
            )
            current_row += 1

        if not rare.empty:
            ws.cell(row=current_row, column=1, value="— Rare combinations below —").font = Font(bold=True)
            current_row += 1
            rare_start = current_row
            for _, item in rare.iterrows():
                _write_spare_row(
                    ws,
                    current_row,
                    item,
                    is_single=is_single,
                    color_percentage=False,
                    single_thresholds=single_thresholds,
                )
                current_row += 1
            ws.row_dimensions.group(rare_start, current_row - 1, hidden=True)

        current_row += 1

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10


def _split_common_and_rare(section: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, bool]:
    common = section[section["attempts"] >= 10].copy()
    rare = section[section["attempts"] < 10].copy()
    limited_data = common.empty and not rare.empty
    if limited_data:
        common = rare[rare["attempts"] > 1].copy()
        rare = rare[rare["attempts"] == 1].copy()
    sort_columns = ["pct", "attempts", "pin_leave"]
    common = common.sort_values(sort_columns, ascending=[True, False, True])
    rare = rare.sort_values(sort_columns, ascending=[True, False, True])
    return common, rare, limited_data


def _write_spare_row(
    ws,
    row: int,
    item: pd.Series,
    *,
    is_single: bool,
    color_percentage: bool,
    single_thresholds: tuple[float, float],
) -> None:
    pin_cell = ws.cell(row=row, column=1, value=str(item["pin_leave"]))
    pin_cell.number_format = "@"
    ws.cell(row=row, column=2, value=int(item["made"]))
    ws.cell(row=row, column=3, value=int(item["attempts"]))
    pct_cell = ws.cell(
        row=row,
        column=4,
        value=float(item["pct"]) if item["pct"] is not None else None,
    )
    pct_cell.number_format = PCT_FORMAT
    if not color_percentage:
        return
    if is_single:
        _apply_percentage_color(
            pct_cell,
            green_threshold=single_thresholds[0],
            yellow_threshold=single_thresholds[1],
        )
    else:
        _apply_percentage_color(
            pct_cell,
            green_threshold=MULTI_PIN_THRESHOLDS[0],
            yellow_threshold=MULTI_PIN_THRESHOLDS[1],
        )


def _aggregate_section(records: pd.DataFrame, is_single: bool) -> pd.DataFrame:
    section = records[records["is_single"] == is_single].copy()
    if section.empty:
        return pd.DataFrame(columns=["pin_leave", "made", "attempts", "pct"])

    section["effective_made"] = (section["made"] & section["count_attempt"]).astype(int)
    section["count_attempt"] = section["count_attempt"].astype(int)

    grouped = (
        section.groupby("pin_leave", dropna=False)
        .agg(
            made=("effective_made", "sum"),
            attempts=("count_attempt", "sum"),
        )
        .reset_index()
    )

    grouped["pct"] = grouped.apply(
        lambda r: (r["made"] / r["attempts"]) if r["attempts"] else None,
        axis=1,
    )

    grouped = grouped[grouped["attempts"] > 0].copy()
    grouped = grouped.sort_values(["attempts", "pin_leave"], ascending=[False, True])
    return grouped


def _all_athletes(
    records: pd.DataFrame,
    frame_attempts: pd.DataFrame | None,
    baker_frames: pd.DataFrame | None,
    team_games: pd.DataFrame | None,
) -> list[str]:
    athletes = set(records["athlete"].dropna().astype(str)) if not records.empty else set()
    for frame in (frame_attempts, baker_frames, team_games):
        if frame is not None and not frame.empty:
            athletes.update(frame["athlete"].dropna().astype(str))
    return sorted(athletes)


def _apply_percentage_color(cell, *, green_threshold: float, yellow_threshold: float) -> None:
    if cell.value is None:
        return

    if cell.value >= green_threshold:
        cell.fill = GREEN_FILL
    elif cell.value >= yellow_threshold:
        cell.fill = YELLOW_FILL
    else:
        cell.fill = RED_FILL


def _unique_sheet_name(name: str, used: set[str]) -> str:
    base = (name or "Athlete").strip()[:31] or "Athlete"
    if base not in used:
        return base

    i = 1
    while True:
        suffix = f"_{i}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        if candidate not in used:
            return candidate
        i += 1
