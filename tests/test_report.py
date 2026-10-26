from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.report import (
    add_team_game_score_columns,
    append_event_total_row,
    build_report_workbook,
    build_summary_workbook,
)


def test_summary_formats_baker_average_as_decimal_points() -> None:
    records = pd.DataFrame(
        [
            {
                "athlete": "Alex",
                "is_single": True,
                "count_attempt": True,
                "made": True,
                "pin_leave": "10",
            }
        ]
    )
    baker_frames = pd.DataFrame(
        [
            {"athlete": "Alex", "points": 19},
            {"athlete": "Alex", "points": 16},
        ]
    )

    workbook_bytes = build_report_workbook(records, baker_frames=baker_frames)
    workbook = load_workbook(BytesIO(workbook_bytes))
    summary = workbook["Overall Summary"]

    assert summary["D2"].number_format == "0.00%"
    assert summary["M2"].value == 17.5
    assert summary["M2"].number_format == "0.0"


def test_summary_uses_distinct_single_and_multi_pin_color_thresholds() -> None:
    records = pd.DataFrame(
        [
            {
                "athlete": "Alex",
                "is_single": True,
                "count_attempt": True,
                "made": attempt < 17,
                "pin_leave": "10",
            }
            for attempt in range(20)
        ]
        + [
            {
                "athlete": "Alex",
                "is_single": False,
                "count_attempt": True,
                "made": attempt < 12,
                "pin_leave": "2-4-5",
            }
            for attempt in range(20)
        ]
    )

    workbook = load_workbook(BytesIO(build_report_workbook(records)))
    summary = workbook["Overall Summary"]

    assert summary["D2"].value == 0.85
    assert summary["D2"].fill.fgColor.rgb == "00FFEB9C"
    assert summary["G2"].value == 0.6
    assert summary["G2"].fill.fgColor.rgb == "00FFEB9C"

    women_workbook = load_workbook(BytesIO(build_report_workbook(records, team="Women")))
    women_summary = women_workbook["Overall Summary"]
    women_athlete = women_workbook["Alex"]

    assert women_summary["D2"].value == 0.85
    assert women_summary["D2"].fill.fgColor.rgb == "00C6EFCE"
    assert women_athlete["D5"].value == 0.85
    assert women_athlete["D5"].fill.fgColor.rgb == "00C6EFCE"


def test_summary_average_uses_team_game_scores() -> None:
    records = pd.DataFrame(
        [
            {
                "athlete": "Alex",
                "is_single": True,
                "count_attempt": True,
                "made": True,
                "pin_leave": "10",
            }
        ]
    )
    team_games = pd.DataFrame(
        [
            {"athlete": "Alex", "score": 200},
            {"athlete": "Alex", "score": 180},
        ]
    )

    workbook = load_workbook(BytesIO(build_report_workbook(records, team_games=team_games)))
    summary = workbook["Overall Summary"]

    assert summary["N1"].value == "Average"
    assert summary["N2"].value == 190
    assert summary["N2"].number_format == "0.00"


def test_athlete_sheet_uses_limited_data_common_and_one_attempt_rare() -> None:
    records = pd.DataFrame(
        [
            {
                "athlete": "Alex",
                "is_single": True,
                "count_attempt": True,
                "made": made,
                "pin_leave": "10",
            }
            for made in (True, False)
        ]
        + [
            {
                "athlete": "Alex",
                "is_single": True,
                "count_attempt": True,
                "made": False,
                "pin_leave": "7",
            }
        ]
    )

    workbook = load_workbook(BytesIO(build_report_workbook(records)))
    athlete = workbook["Alex"]

    assert athlete["A3"].value == "Single Pin"
    assert athlete["A5"].value == "10"
    assert athlete["C5"].value == 2
    assert athlete["A6"].value == "— Rare combinations below —"
    assert athlete["A7"].value == "7"
    assert athlete["C7"].value == 1
    assert athlete.row_dimensions[7].hidden is True
    assert athlete["A9"].value == "Multi Pin"


def test_filtered_summary_workbook_contains_only_supplied_summary() -> None:
    summary = pd.DataFrame(
        [
            {
                "Name": "Alex",
                "Single-Pin Made": 9,
                "Single-Pin Attempts": 10,
                "Single-Pin %": 0.9,
                "Multi-Pin Made": 4,
                "Multi-Pin Attempts": 5,
                "Multi-Pin %": 0.8,
                "Strikes": 20,
                "Strike Attempts": 40,
                "Strike %": 0.5,
                "Baker Pts": 180,
                "Baker Frames": 10,
                "Baker Avg": 18.0,
                "Average": 205.25,
            }
        ]
    )

    workbook = load_workbook(BytesIO(build_summary_workbook(summary, "Men", "2026 WHAC 1")))
    worksheet = workbook["Overall Summary"]

    assert workbook.sheetnames == ["Overall Summary"]
    assert worksheet["A1"].value == "LTU Men - 2026 WHAC 1 Overall Summary"
    assert worksheet["A3"].value == "Name"
    assert worksheet["D3"].value == "SP %"
    assert worksheet["A4"].value == "Alex"
    assert worksheet["D4"].value == 0.9
    assert worksheet["D4"].number_format == "0.00%"
    assert worksheet["N4"].value == 205.25
    assert worksheet.freeze_panes == "A4"


def test_team_game_scores_are_chronological_and_inserted_before_average() -> None:
    summary = pd.DataFrame(
        [
            {"Name": "Alex", "Baker Avg": 18.2, "Average": 205.0},
            {"Name": "Blake", "Baker Avg": 17.8, "Average": 190.0},
        ]
    )
    team_games = pd.DataFrame(
        [
            {"athlete": "Alex", "game_date": "2026-10-02 11:00", "game_id": "a:2", "score": 210},
            {"athlete": "Alex", "game_date": "2026-10-02 09:00", "game_id": "a:1", "score": 200},
            {"athlete": "Blake", "game_date": "2026-10-02 10:00", "game_id": "b:1", "score": 190},
        ]
    )

    result = add_team_game_score_columns(summary, team_games)

    assert result.columns.tolist() == ["Name", "Baker Avg", "Game 1", "Game 2", "Average"]
    assert result.loc[0, ["Game 1", "Game 2"]].tolist() == [200, 210]
    assert result.loc[1, "Game 1"] == 190
    assert pd.isna(result.loc[1, "Game 2"])

    workbook = load_workbook(BytesIO(build_summary_workbook(result, "Men", "2026 WHAC 1")))
    headers = [cell.value for cell in workbook["Overall Summary"][3]]
    assert headers == ["Name", "BKR Avg", "Game 1", "Game 2", "AVG"]
    assert workbook["Overall Summary"]["D5"].value is None


def test_selected_athlete_uses_filter_name_for_event_summary_row() -> None:
    summary = pd.DataFrame([{"Name": "2026 WHAC 1", "Baker Avg": 18.0, "Average": 220.0}])
    team_games = pd.DataFrame(
        [{"athlete": "Alex", "game_date": "2026-10-02", "game_id": "a:1", "score": 220}]
    )

    result = add_team_game_score_columns(summary, team_games, athlete_filter="Alex")

    assert result.loc[0, "Game 1"] == 220


def test_event_total_recalculates_rates_and_weighted_averages() -> None:
    summary = pd.DataFrame(
        [
            {
                "Name": "Alex",
                "Single-Pin Made": 8,
                "Single-Pin Attempts": 10,
                "Single-Pin %": 0.8,
                "Multi-Pin Made": 3,
                "Multi-Pin Attempts": 5,
                "Multi-Pin %": 0.6,
                "Strikes": 20,
                "Strike Attempts": 40,
                "Strike %": 0.5,
                "Baker Pts": 180,
                "Baker Frames": 10,
                "Baker Avg": 18.0,
                "Average": 200.0,
            },
            {
                "Name": "Blake",
                "Single-Pin Made": 1,
                "Single-Pin Attempts": 2,
                "Single-Pin %": 0.5,
                "Multi-Pin Made": 4,
                "Multi-Pin Attempts": 5,
                "Multi-Pin %": 0.8,
                "Strikes": 15,
                "Strike Attempts": 20,
                "Strike %": 0.75,
                "Baker Pts": 85,
                "Baker Frames": 5,
                "Baker Avg": 17.0,
                "Average": 220.0,
            },
        ]
    )
    team_games = pd.DataFrame(
        [
            {"athlete": "Alex", "score": 190},
            {"athlete": "Alex", "score": 210},
            {"athlete": "Blake", "score": 220},
        ]
    )

    result = append_event_total_row(summary, team_games)
    total = result.iloc[-1]

    assert total["Name"] == "Total"
    assert total["Single-Pin Made"] == 9
    assert total["Single-Pin Attempts"] == 12
    assert total["Single-Pin %"] == 0.75
    assert total["Multi-Pin %"] == 0.7
    assert total["Strike %"] == 35 / 60
    assert total["Baker Pts"] == 265
    assert total["Baker Frames"] == 15
    assert total["Baker Avg"] == 265 / 15
    assert total["Average"] == 620 / 3
